# -*- coding: utf-8 -*-
"""
Excel 报告输出：多 Sheet + 条件格式 + 图表。

只依赖 openpyxl。若未安装，自动降级为 CSV 分文件输出（见 write_csv_fallback）。
"""
from __future__ import annotations

import csv
import os
from typing import Any, Dict, List, Sequence

from .models import ItemComparison, Quote, SupplierScore, WEIGHT_LABELS


HEADER_FILL = "1F4E79"
BEST_FILL = "C6EFCE"
WARN_FILL = "FFEB9C"
BAD_FILL = "FFC7CE"


def _try_openpyxl():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.chart import BarChart, Reference
        from openpyxl.formatting.rule import ColorScaleRule
        return Workbook, Alignment, Border, Font, PatternFill, Side, get_column_letter, BarChart, Reference, ColorScaleRule
    except ImportError:
        return None


def _sheet_from_rows(ws, rows: Sequence[Dict[str, Any]],
                     header_fill: str = HEADER_FILL,
                     min_width: int = 10, max_width: int = 42):
    """把 dict 列表写进 sheet，并做基础美化。"""
    packs = _try_openpyxl()
    if packs is None:
        return
    _, Alignment, Border, Font, PatternFill, Side, get_column_letter, *_ = packs
    if not rows:
        ws["A1"] = "（无数据）"
        return

    cols: List[str] = []
    for r in rows:
        for k in r.keys():
            if k not in cols:
                cols.append(k)

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=header_fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for i, r in enumerate(rows, 2):
        for j, c in enumerate(cols, 1):
            v = r.get(c)
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=isinstance(v, str) and len(str(v)) > 30)
            if isinstance(v, float):
                cell.number_format = "#,##0.0000" if abs(v) < 100 else "#,##0.00"

    for j, c in enumerate(cols, 1):
        longest = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows])
        width = min(max(min_width, longest * 1.6 + 2), max_width)
        ws.column_dimensions[get_column_letter(j)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"


def write_report(path: str, context: Dict[str, Any]) -> str:
    """
    context 需要包含：
      summary / matrix_rows(含 suppliers) / detail_rows / score_rows / inventory(dict|None)
    """
    packs = _try_openpyxl()
    if packs is None:
        return write_csv_fallback(path, context)
    Workbook, Alignment, Border, Font, PatternFill, Side, get_column_letter, BarChart, Reference, ColorScaleRule = packs

    wb = Workbook()
    wb.remove(wb.active)

    # ---------- 1. 汇总 ----------
    ws = wb.create_sheet("① 汇总看板")
    summary: Dict[str, Any] = context.get("summary", {})
    rows = [{"指标": k, "数值": v} for k, v in summary.items()]
    ws["A1"] = "MRO 询价比价与供应商评估 · 汇总看板"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A2"] = context.get("generated_at", "")
    ws["A2"].font = Font(size=9, color="808080")
    _sheet_from_rows(ws, rows, header_fill=HEADER_FILL)
    ws.freeze_panes = "A4"

    # ---------- 2. 比价矩阵 ----------
    matrix = context.get("matrix_rows", [])
    ws2 = wb.create_sheet("② 比价矩阵")
    _sheet_from_rows(ws2, matrix)
    _highlight_matrix(ws2, matrix, context.get("suppliers", []), PatternFill, Font)

    # ---------- 3. 报价明细 ----------
    ws3 = wb.create_sheet("③ 报价明细")
    _sheet_from_rows(ws3, context.get("detail_rows", []), header_fill="2E75B6")

    # ---------- 4. 供应商评分 ----------
    score_rows = context.get("score_rows", [])
    ws4 = wb.create_sheet("④ 供应商评分")
    _sheet_from_rows(ws4, score_rows, header_fill="548235")
    if score_rows:
        _add_score_chart(ws4, score_rows, BarChart, Reference)

    # ---------- 5. 库存补货 ----------
    inv = context.get("inventory")
    if inv:
        ws5 = wb.create_sheet("⑤ 库存与补货")
        _sheet_from_rows(ws5, inv.get("rows", []), header_fill="BF8F00")
        start = len(inv.get("rows", [])) + 3
        ws5.cell(row=start, column=1, value="分类汇总").font = Font(bold=True, size=12)
        cls_rows = [{"ABC分类": k, **v} for k, v in inv.get("by_class", {}).items()]
        _sheet_from_rows_at(ws5, cls_rows, start_row=start + 1, header_fill="BF8F00")

    # ---------- 6. 决策建议 ----------
    ws6 = wb.create_sheet("⑥ 决策建议")
    _sheet_from_rows(ws6, context.get("advice_rows", []), header_fill="7030A0")
    ws6.column_dimensions["D"].width = 60

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)
    return path


def _sheet_from_rows_at(ws, rows, start_row: int, header_fill: str):
    packs = _try_openpyxl()
    _, Alignment, Border, Font, PatternFill, Side, get_column_letter, *_ = packs
    if not rows:
        return
    cols: List[str] = []
    for r in rows:
        for k in r.keys():
            if k not in cols:
                cols.append(k)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=start_row, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=header_fill)
    for i, r in enumerate(rows, start_row + 1):
        for j, c in enumerate(cols, 1):
            ws.cell(row=i, column=j, value=r.get(c))


def _highlight_matrix(ws, matrix, suppliers, PatternFill, Font):
    """最优供应商列标绿，风险提示标黄。"""
    if not matrix:
        return
    headers = [c.value for c in ws[1]]
    try:
        best_col = headers.index("最优供应商") + 1
    except ValueError:
        best_col = None
    green = PatternFill("solid", fgColor=BEST_FILL)
    for i in range(2, len(matrix) + 2):
        if best_col:
            ws.cell(row=i, column=best_col).fill = green
    for name, fill, color in (("风险提示", WARN_FILL, None),):
        if name in headers:
            col = headers.index(name) + 1
            for i in range(2, len(matrix) + 2):
                v = ws.cell(row=i, column=col).value
                if v:
                    ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=fill)
    # 价格列做色阶
    packs = _try_openpyxl()
    ColorScaleRule = packs[-1]
    for s in suppliers:
        if s in headers:
            col = headers.index(s) + 1
            from openpyxl.utils import get_column_letter
            rng = f"{get_column_letter(col)}2:{get_column_letter(col)}{len(matrix) + 1}"
            ws.conditional_formatting.add(rng, ColorScaleRule(
                start_type="min", start_color="F8696B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="63BE7B"))


def _add_score_chart(ws, score_rows, BarChart, Reference):
    headers = [c.value for c in ws[1]]
    if "供应商" not in headers or "综合得分" not in headers:
        return
    n = len(score_rows)
    chart = BarChart()
    chart.type = "col"
    chart.title = "供应商综合得分对比"
    chart.y_axis.title = "综合得分"
    chart.height = 8
    chart.width = 16
    s_col = headers.index("供应商") + 1
    v_col = headers.index("综合得分") + 1
    data = Reference(ws, min_col=v_col, min_row=1, max_row=n + 1)
    cats = Reference(ws, min_col=s_col, min_row=2, max_row=n + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"A{n + 4}")


def write_csv_fallback(path: str, context: Dict[str, Any]) -> str:
    """openpyxl 不可用时的降级方案：按区块输出多个 CSV。"""
    base = os.path.splitext(path)[0]
    os.makedirs(os.path.dirname(os.path.abspath(base)) or ".", exist_ok=True)
    blocks = {
        "_summary.csv": context.get("summary_rows", []),
        "_matrix.csv": context.get("matrix_rows", []),
        "_detail.csv": context.get("detail_rows", []),
        "_score.csv": context.get("score_rows", []),
    }
    for suffix, rows in blocks.items():
        if not rows:
            continue
        cols: List[str] = []
        for r in rows:
            for k in r.keys():
                if k not in cols:
                    cols.append(k)
        with open(base + suffix, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
    return base + "_summary.csv"


def build_context(comparisons, scores, weights, inv_result=None, generated_at: str = "") -> Dict[str, Any]:
    """把引擎结果整理成写报告需要的上下文。"""
    from .compare import build_matrix, summary as cmp_summary

    detail_rows: List[Dict] = []
    for c in comparisons:
        for q in c.quotes:
            r = {
                "物料编码": q.item_code,
                "物料名称": q.item_name,
                "物料类别": q.category,
                "规格型号": q.spec,
                "单位": q.unit,
                "需求数量": q.qty,
                "供应商": q.supplier,
                "报价单价": round(q.unit_price, 4),
                "税率": q.tax_rate,
                "是否含税": "是" if q.tax_included else "否",
                "不含税单价": round(q.net_unit_price, 4),
                "含税单价": round(q.gross_unit_price, 4),
                "实际下单量": q.order_qty,
                "运费(含税)": q.freight,
                "到货总成本(不含税)": round(q.landed_cost, 2),
                "到货等效单价(不含税)": round(q.landed_unit_price, 4),
                "到货等效单价(含税)": round(q.landed_unit_price_gross, 4),
                "交期(天)": q.lead_time,
                "起订量": q.moq,
                "付款条件": q.payment_terms,
                "资质合规": "是" if q.cert_ok else "否",
                "历史单价": q.history_price or None,
                "降本率": round(q.saving_rate, 4) if q.saving_rate is not None else None,
                "是否最优": "★最优" if c.best and q.supplier == c.best.supplier and q is c.best else "",
            }
            detail_rows.append(r)

    matrix = build_matrix(comparisons)
    sum_stats = cmp_summary(comparisons)

    advice_rows: List[Dict] = []
    for c in comparisons:
        advice_rows.append({
            "物料编码": c.item_code,
            "物料名称": c.item_name,
            "报价家数": c.supplier_count,
            "最优供应商": c.best.supplier if c.best else "",
            "最优到货单价(不含税)": round(c.best.landed_unit_price, 4) if c.best else None,
            "价差率": round(c.spread_rate, 4),
            "与次低价差": round(c.gap_to_second, 4),
            "比最高价节约": round(c.saving_vs_worst, 2),
            "相对历史价节约": round(c.saving_vs_history, 2),
            "风险提示": "；".join(c.risk_flags()),
            "建议": c.recommendation(),
        })

    from .scoring import to_rows
    score_rows = to_rows(scores)
    weight_rows = [{"维度": WEIGHT_LABELS.get(k, k), "权重": v} for k, v in weights.items()]

    return {
        "summary": {**sum_stats, **{f"权重-{WEIGHT_LABELS.get(k, k)}": v for k, v in weights.items()}},
        "summary_rows": [{"指标": k, "数值": v} for k, v in sum_stats.items()] + weight_rows,
        "matrix_rows": matrix["rows"],
        "suppliers": matrix["suppliers"],
        "detail_rows": detail_rows,
        "score_rows": score_rows,
        "advice_rows": advice_rows,
        "inventory": inv_result,
        "generated_at": generated_at,
    }
